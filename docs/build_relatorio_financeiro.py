#!/usr/bin/env python3
"""Relatório financeiro (.docx) — posição atual + projeção de 3 meses (MODELO).
Valores financeiros são ILUSTRATIVOS; as seções de sistemas/estrutura/passos
descrevem o trabalho real. Identidade visual setup.com.br."""
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

A = "/home/user/rstoi/docs/assets"
OUT = "/home/user/rstoi/docs/Relatorio-Financeiro-setup.com.br.docx"

INK="003C54"; BLUE="0E6E92"; TEAL="1F97B4"; STEEL="609CC0"; AMBER="4E86AE"
SLATE="5E6A72"; BODY="212B36"; GRAY="A8A8A8"; LIGHT="E4EEF4"; LIGHT2="E1F0F4"
WARN="FBEFD6"; WHITE="FFFFFF"
def C(h): return RGBColor.from_string(h)

# ---------------------------------------------------------------- dados (ILUSTRATIVOS)
meses = ["Mês +1", "Mês +2", "Mês +3"]
receita  = [268, 276, 285]      # R$ mil
despesas = [208, 211, 214]
resultado= [r-d for r,d in zip(receita,despesas)]
caixa0   = 480
caixa    = []
acc = caixa0
for r in resultado:
    acc += r; caixa.append(acc)

# ---------------------------------------------------------------- gráfico projeção
def chart():
    fig, ax = plt.subplots(figsize=(8.6, 3.9), dpi=200)
    x = np.arange(len(meses)); w = 0.34
    ax.bar(x-w/2, receita,  w, label="Receita", color="#0E6E92")
    ax.bar(x+w/2, despesas, w, label="Despesas", color="#609CC0")
    ax.plot(x, caixa, "-o", color="#1F97B4", lw=2.4, label="Caixa acumulado")
    for i,v in enumerate(caixa):
        ax.annotate(f"{v}", (x[i], caixa[i]), textcoords="offset points",
                    xytext=(0,8), ha="center", fontsize=8, color="#003C54", weight="bold")
    ax.set_xticks(x); ax.set_xticklabels(meses, fontsize=9, color="#003C54")
    ax.set_ylabel("R$ mil", fontsize=9, color="#5E6A72")
    ax.tick_params(colors="#5E6A72")
    for s in ["top","right"]: ax.spines[s].set_visible(False)
    ax.legend(frameon=False, fontsize=8.5, ncol=3, loc="upper center", bbox_to_anchor=(0.5,1.16))
    ax.text(0.99,0.02,"EXEMPLO ILUSTRATIVO", transform=ax.transAxes, ha="right",
            color="#A8A8A8", fontsize=20, weight="bold", alpha=0.25)
    fig.tight_layout()
    fig.savefig(f"{A}/fin_projecao.png", bbox_inches="tight", facecolor="white")
    plt.close(fig)
chart()

# ---------------------------------------------------------------- helpers docx
doc = Document()
for s in doc.sections:
    s.top_margin=Inches(0.85); s.bottom_margin=Inches(0.85)
    s.left_margin=Inches(0.9); s.right_margin=Inches(0.9)
st = doc.styles["Normal"]; st.font.name="Calibri"; st.font.size=Pt(11)
st.font.color.rgb=C(BODY); st.paragraph_format.space_after=Pt(6); st.paragraph_format.line_spacing=1.15

def cell_bg(cell,h):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd")
    shd.set(qn("w:val"),"clear"); shd.set(qn("w:fill"),h); tcPr.append(shd)
def cell_marg(cell,t=70,b=70,l=110,r=110):
    tcPr=cell._tc.get_or_add_tcPr(); m=OxmlElement("w:tcMar")
    for tag,v in (("top",t),("bottom",b),("start",l),("end",r)):
        e=OxmlElement(f"w:{tag}"); e.set(qn("w:w"),str(v)); e.set(qn("w:type"),"dxa"); m.append(e)
    tcPr.append(m)
def no_borders(tbl):
    el=OxmlElement("w:tblBorders")
    for edge in ("top","left","bottom","right","insideH","insideV"):
        e=OxmlElement(f"w:{edge}"); e.set(qn("w:val"),"none"); el.append(e)
    tbl._tbl.tblPr.append(el)
def run(p,t,size=11,color=BODY,bold=False,italic=False):
    r=p.add_run(t); r.font.name="Calibri"; r.font.size=Pt(size); r.font.color.rgb=C(color)
    r.bold=bold; r.italic=italic; return r
def heading(num,title):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(15); p.paragraph_format.space_after=Pt(6)
    if num: run(p,f"{num}  ",15,TEAL,True)
    run(p,title,15,INK,True)
    pPr=p._p.get_or_add_pPr(); b=OxmlElement("w:pBdr"); bot=OxmlElement("w:bottom")
    bot.set(qn("w:val"),"single"); bot.set(qn("w:sz"),"6"); bot.set(qn("w:space"),"4"); bot.set(qn("w:color"),"D2E2EC")
    b.append(bot); pPr.append(b)
def body(t,italic=False,color=BODY):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); run(p,t,11,color,italic=italic)
def bullet(prefix,t):
    p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(3)
    if prefix: run(p,prefix,11,INK,True)
    run(p,t,11,BODY)
def callout(label,t,bg=WARN):
    tb=doc.add_table(rows=1,cols=1); no_borders(tb); c=tb.rows[0].cells[0]
    cell_bg(c,bg); cell_marg(c,130,130,180,180); p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    run(p,label,11,INK,True); run(p,t,11,BODY)
def table(headers,rows,widths,total_row=None):
    n=len(rows)+1+(1 if total_row else 0)
    tb=doc.add_table(rows=n,cols=len(headers)); tb.alignment=WD_TABLE_ALIGNMENT.CENTER; no_borders(tb)
    for j,h in enumerate(headers):
        c=tb.rows[0].cells[j]; c.width=Inches(widths[j]); cell_bg(c,INK); cell_marg(c)
        p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); run(p,h,10,WHITE,True)
    for i,rowv in enumerate(rows):
        for j,val in enumerate(rowv):
            c=tb.rows[i+1].cells[j]; c.width=Inches(widths[j]); cell_bg(c, "F1F4F6" if i%2 else WHITE); cell_marg(c)
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
            run(p,val,10,(INK if j==0 else BODY),bold=(j==0))
    if total_row:
        for j,val in enumerate(total_row):
            c=tb.rows[-1].cells[j]; c.width=Inches(widths[j]); cell_bg(c,LIGHT2); cell_marg(c)
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); run(p,val,10,INK,True)
def image(path,w=6.4,cap=None):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(2)
    p.add_run().add_picture(path,width=Inches(w))
    if cap:
        c=doc.add_paragraph(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after=Pt(10)
        run(c,cap,9,SLATE,italic=True)

# ---------------------------------------------------------------- CAPA
doc.add_picture(f"{A}/00_capa.png", width=Inches(6.9))
doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
sp=doc.add_paragraph(); sp.paragraph_format.space_after=Pt(10)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
run(p,"Relatório Financeiro",18,INK,True)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(2)
run(p,"Posição atual e projeção de 3 meses",13,BODY)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
run(p,"Junho / 2026  ·  Uso interno  ·  MODELO (valores ilustrativos)",10.5,SLATE,italic=True)
sp=doc.add_paragraph(); sp.paragraph_format.space_after=Pt(14)
callout("⚠️  Aviso importante:  ",
        "este relatório é um MODELO. Nenhum sistema financeiro real foi consultado nesta "
        "geração — os valores monetários são ILUSTRATIVOS e servem de gabarito. Substitua-os "
        "pelos dados reais ao conectar os sistemas de contratos, comercial e financeiro "
        "(ver seções 3 e 6).")
doc.add_page_break()

# ---------------------------------------------------------------- 1. POSIÇÃO FINANCEIRA
heading("1","Posição financeira atual  (ilustrativa)")
body("Fotografia do momento. Os valores abaixo são exemplos para validar o formato; "
     "ao conectar o sistema financeiro/ERP, esta tabela é preenchida automaticamente.")
table(["Indicador","Valor (R$ mil) — ILUSTRATIVO","Observação"],
      [["Caixa e equivalentes","480","saldo disponível"],
       ["Contas a receber","320","carteira 30/60/90 dias"],
       ["Contas a pagar","210","obrigações do período"],
       ["Capital de giro líquido","590","receber + caixa − pagar"],
       ["Receita recorrente (MRR)","260","base de contratos ativos"],
       ["Despesas operacionais / mês","205","custo + despesa fixa/variável"]],
      [2.3,2.3,2.0],
      total_row=["Resultado operacional / mês","55","MRR − despesas (ilustrativo)"])
body("")
callout("Como ler:  ","capital de giro positivo e resultado mensal positivo indicam folga de "
        "curto prazo — confirme com os números reais antes de qualquer decisão.", bg=LIGHT2)

# ---------------------------------------------------------------- 2. PROJEÇÃO
heading("2","Projeção de 3 meses  (ilustrativa)")
body("Cenário-base com crescimento moderado de receita (~3%/mês) e despesas (~1,5%/mês). "
     "Premissas explícitas para você ajustar:")
bullet("Receita: ","crescimento de ~3%/mês sobre o MRR atual (novos contratos − churn).")
bullet("Despesas: ","crescimento de ~1,5%/mês (inflação de custos + contratações).")
bullet("Caixa acumulado: ","caixa inicial + soma dos resultados mensais.")
image(f"{A}/fin_projecao.png", 6.2, "Figura 1 — Projeção de receita, despesas e caixa acumulado (ilustrativa).")
table(["Período","Receita","Despesas","Resultado","Caixa acumulado"],
      [[meses[i], str(receita[i]), str(despesas[i]), str(resultado[i]), str(caixa[i])] for i in range(3)],
      [1.7,1.4,1.4,1.4,1.7],
      total_row=["Total 3 meses", str(sum(receita)), str(sum(despesas)), str(sum(resultado)), f"{caixa[-1]} (final)"])
body("Valores em R$ mil. ILUSTRATIVO — substituir por forecast do comercial (pipeline) e "
     "agenda de recebíveis (contratos).", italic=True)
doc.add_page_break()

# ---------------------------------------------------------------- 3. DADOS E SISTEMAS
heading("3","Dados e sistemas acessados")
body("Transparência sobre as fontes. Status real nesta sessão:")
table(["Sistema","Status","O que forneceria ao relatório"],
      [["Repositório GitHub (rstoi/rstoi)","✅ acessado","código, scripts e este relatório"],
       ["Google Workspace (Drive/Gmail/Calendar)","conectores disponíveis, NÃO consultados","planilhas/recebíveis, e-mails de cobrança, agenda"],
       ["WhatsApp Business (MCP)","construído, NÃO operacional (rede bloqueada)","cobranças/leads via mensagens"],
       ["computer-use (MCP)","✅ operacional no container","operar sistemas legados sem API"],
       ["Sistema de Contratos (interno)","NÃO conectado","MRR, recebíveis, renovações"],
       ["Sistema Comercial (interno)","NÃO conectado","pipeline → forecast de receita"],
       ["ERP / Financeiro","NÃO conectado","caixa, contas a pagar/receber, DRE"]],
      [2.4,2.4,1.8])
body("")
callout("Conclusão honesta:  ","nenhuma fonte financeira real foi consultada nesta geração — "
        "por isso os números das seções 1 e 2 são ilustrativos. Conectados os sistemas acima, "
        "o relatório passa a refletir dados reais.")

# ---------------------------------------------------------------- 4. ESTRUTURA CRIADA
heading("4","Estrutura criada")
body("O que de fato foi construído e versionado no repositório (base para automatizar este relatório):")
bullet("MCP WhatsApp Business — ","29 ferramentas (mensagens, grupos, contatos, mídia) + guardrail de grupos bloqueados.")
bullet("Agente /setup — ","operações por mensagem, autorização por participação em grupo (deny por padrão).")
bullet("MCP computer-use — ","16 ferramentas (tela, mouse, teclado, run_command) para operar sistemas sem API.")
bullet("Conectores Google Workspace — ","Gmail, Calendar e Drive disponíveis via MCP.")
bullet("Resiliência (SessionStart hook) — ","reconstrói o ambiente a cada sessão (npm/pip/Chromium).")
bullet("Documentação executiva — ","documento, deck e PDFs com a identidade da marca + scripts reprodutíveis.")

# ---------------------------------------------------------------- 5. PASSOS
heading("5","Passos para a criação da estrutura")
for i,(t,d) in enumerate([
    ("Documento e identidade","relatório/deck executivos e extração da paleta + logo da marca."),
    ("Resiliência","SessionStart hook para reinstalar dependências a cada reboot."),
    ("Guardrail","WA_BLOCKED_GROUPS para nunca monitorar/interagir com o grupo financasfacil."),
    ("Ativação do MCP","habilitar o servidor de WhatsApp + smoke test funcional (round-trip)."),
    ("Agente seguro","escopo de grupos (financeiro setup, projetos setup) + autorização por membro."),
    ("Bootstrap de rede","script que valida o egresso e conduz connect/agent após liberação."),
    ("Consolidação","merge do PR #1 no branch padrão; tudo testado (48/48)."),
],1):
    bullet(f"Passo {i} — {t}: ", d)

# ---------------------------------------------------------------- 6. COMPLEMENTOS
heading("6","Complementos e próximos passos")
body("Para transformar este modelo em relatório com dados reais:")
bullet("Conectar fontes — ","ERP/financeiro (caixa, contas), Contratos (MRR/recebíveis), Comercial (pipeline→forecast).")
bullet("Automatizar a coleta — ","agente lê as planilhas/relatórios no Drive ou via API e preenche as tabelas.")
bullet("Atualização periódica — ","gerar o .docx semanalmente/mensalmente por script (como este).")
bullet("Indicadores a acompanhar — ","MRR, churn, runway (meses de caixa), DSO (prazo médio de recebimento), margem.")
bullet("Controles — ","fonte e data de cada número no rodapé; aprovação humana antes de divulgar.")
callout("Como gerar com dados reais:  ","substitua as listas de valores no topo de "
        "docs/build_relatorio_financeiro.py (ou aponte o script para a fonte de dados) e rode "
        "novamente — o layout e os gráficos se atualizam sozinhos.", bg=LIGHT)

# ---------------------------------------------------------------- rodapé
def footer():
    for section in doc.sections:
        f=section.footer; lp=f.paragraphs[0]; lp.alignment=WD_ALIGN_PARAGRAPH.LEFT
        lp.paragraph_format.space_after=Pt(0)
        lp.add_run().add_picture(f"{A}/setup_logo@hi.png", height=Inches(0.16))
        p=f.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        run(p,"setup.com.br · Relatório Financeiro (MODELO) · valores ilustrativos",8,SLATE)
footer()

doc.save(OUT)
print("Relatório salvo em", OUT)
