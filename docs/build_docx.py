#!/usr/bin/env python3
"""Monta o documento executivo .docx com identidade visual e gráficos."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

A = "/home/user/rstoi/docs/assets"
OUT = "/home/user/rstoi/docs/Infraestrutura-IA-setup.com.br.docx"

# Paleta extraída do logo oficial setup.com.br
INK   = RGBColor(0x00, 0x3C, 0x54)   # petróleo escuro ("set")
BLUE  = RGBColor(0x0E, 0x6E, 0x92)   # azul setup (primário)
TEAL  = RGBColor(0x1F, 0x97, 0xB4)   # teal do swoosh
AMBER = RGBColor(0x4E, 0x86, 0xAE)   # azul aço médio (4º tom)
STEEL = RGBColor(0x60, 0x9C, 0xC0)   # azul aço claro
SLATE = RGBColor(0x5E, 0x6A, 0x72)   # cinza-azulado (texto secundário)
BODY  = RGBColor(0x21, 0x2B, 0x36)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

doc = Document()

# ---- margens ----
for s in doc.sections:
    s.top_margin = Inches(0.85)
    s.bottom_margin = Inches(0.85)
    s.left_margin = Inches(0.9)
    s.right_margin = Inches(0.9)

# ---- estilo base ----
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = BODY
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15


def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        e = OxmlElement(f"w:{tag}")
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tcPr.append(m)


def no_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "none")
        borders.append(e)
    tblPr.append(borders)


def shade_para(p, hex_color):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


def add_run(p, text, size=11, color=BODY, bold=False, italic=False, font="Calibri"):
    r = p.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.bold = bold
    r.italic = italic
    return r


def heading(text, num=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    if num:
        add_run(p, f"{num}  ", size=15, color=TEAL, bold=True)
    add_run(p, text, size=15, color=INK, bold=True)
    # linha sob o título
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "D2E2EC")
    pbdr.append(bottom)
    pPr.append(pbdr)
    return p


def subheading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    add_run(p, text, size=12, color=BLUE, bold=True)
    return p


def body(text, italic=False, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    add_run(p, text, size=11, color=BODY, italic=italic)
    return p


def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    if bold_prefix:
        add_run(p, bold_prefix, size=11, color=INK, bold=True)
    add_run(p, text, size=11, color=BODY)
    return p


def add_image(path, width_in=6.6, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Inches(width_in))
    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(10)
        add_run(c, caption, size=9, color=SLATE, italic=True)


# ============================================================ CAPA
doc.add_picture(f"{A}/00_capa.png", width=Inches(6.9))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(14)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_run(p, "Documento Executivo", size=13, color=SLATE, bold=True)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
add_run(p, "Arquitetura de produtividade aumentada por Inteligência Artificial",
        size=12, color=BODY)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_run(p, "Integração PC · Notebook · Tablet · Smartphone · WhatsApp · Google Workspace",
        size=11, color=SLATE, italic=True)

sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(40)

# bloco de metadados (tabela 1x3)
meta = doc.add_table(rows=1, cols=3)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
no_table_borders(meta)
items = [("VERSÃO", "1.0"), ("DATA", "Junho / 2026"), ("CLASSIFICAÇÃO", "Uso interno")]
for cell, (k, v) in zip(meta.rows[0].cells, items):
    set_cell_bg(cell, "E4EEF4")
    set_cell_margins(cell)
    cell.width = Inches(2.2)
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    add_run(p, k + "\n", size=8.5, color=SLATE, bold=True)
    add_run(p, v, size=11, color=INK, bold=True)

doc.add_page_break()

# ============================================================ SUMÁRIO EXECUTIVO
heading("Sumário Executivo", "")
body("A setup.com.br adota uma infraestrutura de TI em que a Inteligência "
     "Artificial deixa de ser uma ferramenta avulsa e passa a ser uma camada "
     "transversal de trabalho. Cada profissional ganha um assistente único — "
     "acessível no notebook, no celular, no tablet e pelo WhatsApp — que entende "
     "o contexto da empresa e executa tarefas de ponta a ponta.")
body("O diferencial não é “conversar com um chatbot”, e sim ter agentes que "
     "operam de verdade: leem e-mails, organizam a agenda, geram documentos no "
     "Drive, respondem clientes no WhatsApp e atuam sobre os sistemas internos já "
     "existentes (projetos, contratos e comercial) — sempre sob controle humano "
     "nas decisões críticas. Para isso combinamos três pilares:")
bullet("seleciona automaticamente o modelo certo para cada tarefa, equilibrando custo, velocidade e qualidade.", "Multimodelo — ")
bullet("o ciclo planejar → agir → observar → corrigir que faz o agente concluir o trabalho sozinho.", "Harness loop — ")
bullet("um padrão único que conecta a IA aos sistemas reais (WhatsApp, Google Workspace, projetos, contratos e comercial).", "Conectores MCP — ")

# caixa de destaque (benefício)
t = doc.add_table(rows=1, cols=1); no_table_borders(t)
cell = t.rows[0].cells[0]; set_cell_bg(cell, "E1F0F4"); set_cell_margins(cell, 140, 140, 200, 200)
p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
add_run(p, "Resultado esperado:  ", size=11, color=INK, bold=True)
add_run(p, "menos tempo em tarefas repetitivas e mais qualidade nos entregáveis — "
        "e, no mesmo movimento, equipes menos sobrecarregadas, com mais segurança "
        "psicológica e melhores condições de trabalho, sob governança e custo "
        "controlados.", size=11, color=BODY)

add_image(f"{A}/01_camadas.png", 6.4, "Figura 1 — A infraestrutura em quatro camadas integradas.")

doc.add_page_break()

# ============================================================ 1. EXPERIÊNCIA
heading("Um assistente em cada tela", "1")
body("Independentemente do dispositivo, o profissional fala com o mesmo agente, "
     "que carrega o mesmo contexto. Pedir pelo WhatsApp “remarca a reunião com o "
     "cliente X para quinta” usa exatamente os mesmos conectores (Calendar + "
     "Gmail) que seriam usados no notebook.")

# tabela de dispositivos
dev = doc.add_table(rows=6, cols=2)
dev.alignment = WD_TABLE_ALIGNMENT.CENTER
no_table_borders(dev)
rows = [
    ("Dispositivo", "Como é usado no dia a dia"),
    ("PC / Notebook", "Assistente de produção: documentos, planilhas, código, análise de dados e operação do computador (computer-use)."),
    ("Tablet", "Revisão, aprovação e leitura de resumos; ditado de tarefas em reunião ou em campo."),
    ("Smartphone", "Captura rápida (foto de nota fiscal, áudio de ideia), consulta de agenda e aprovação de ações."),
    ("WhatsApp", "Canal mais natural no Brasil: o profissional conversa por texto ou áudio; clientes também são atendidos por agentes."),
    ("Chat web / IDE", "Trabalho técnico e suporte: copiloto interno para as equipes."),
]
for i, (a_, b_) in enumerate(rows):
    c0, c1 = dev.rows[i].cells
    c0.width = Inches(1.7); c1.width = Inches(4.9)
    set_cell_margins(c0); set_cell_margins(c1)
    if i == 0:
        set_cell_bg(c0, "003C54"); set_cell_bg(c1, "003C54")
        for cc, tx in ((c0, a_), (c1, b_)):
            p = cc.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            add_run(p, tx, size=10.5, color=WHITE, bold=True)
    else:
        bg = "F1F4F6" if i % 2 else "FFFFFF"
        set_cell_bg(c0, bg); set_cell_bg(c1, bg)
        p = c0.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        add_run(p, a_, size=10.5, color=BLUE, bold=True)
        p = c1.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        add_run(p, b_, size=10.5, color=BODY)

body("")
body("Ponto-chave: é o mesmo cérebro com o mesmo contexto em qualquer tela. "
     "O profissional escolhe o canal mais conveniente; a inteligência e os dados "
     "permanecem consistentes.", italic=True)

# ============================================================ 2. MULTIMODELO
heading("Multimodelo — a IA certa para cada tarefa", "2")
body("Nem toda tarefa exige o modelo mais caro. Um roteador inteligente escolhe "
     "automaticamente entre modelos rápidos e econômicos para o volume do dia a "
     "dia e modelos de máxima capacidade para o trabalho complexo.")
add_image(f"{A}/03_modelos.png", 6.0, "Figura 2 — Roteamento por custo e capacidade.")
body("Além dos modelos de texto, a infraestrutura usa modelos especializados de "
     "transcrição de áudio (áudios de WhatsApp e reuniões), de visão/OCR "
     "(notas fiscais e contratos escaneados) e de busca semântica para consultar "
     "a base de conhecimento da empresa.")

doc.add_page_break()

# ============================================================ 3. HARNESS LOOP
heading("Harness loop — de chatbot a agente que entrega", "3")
body("Um chatbot responde; um agente opera. O harness loop é o ciclo que permite "
     "ao agente concluir uma tarefa de ponta a ponta, corrigindo-se a cada passo.")
add_image(f"{A}/02_loop.png", 5.6, "Figura 3 — O ciclo planejar → agir → observar → refletir.")

subheading("Exemplo real — “Fechar proposta para o cliente X”")
bullet("o vendedor pede pelo WhatsApp; o agente entende o objetivo.", "Planejar — ")
bullet("busca o template e o histórico no Drive, recupera o último acordo de preço no Gmail, gera a proposta, cria o arquivo e agenda a apresentação no Calendar.", "Agir / Observar (loop) — ")
bullet("envia o resumo e o link ao vendedor pelo WhatsApp para aprovação; se ele pedir ajustes, o loop reabre e corrige.", "Aprovação humana — ")
body("Tudo isso sem o profissional sair do WhatsApp. O loop, os modelos e os "
     "conectores estão na infraestrutura — o humano apenas decide e aprova.",
     italic=True)

# caixa governança no loop
t = doc.add_table(rows=1, cols=1); no_table_borders(t)
cell = t.rows[0].cells[0]; set_cell_bg(cell, "E8F1F8"); set_cell_margins(cell, 140, 140, 200, 200)
p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
add_run(p, "Controle humano (human-in-the-loop):  ", size=11, color=INK, bold=True)
add_run(p, "ações de baixo risco (rascunhar, resumir, buscar) são automáticas; "
        "ações externas ou irreversíveis (enviar e-mail ao cliente, apagar "
        "arquivo, alterar contrato) exigem aprovação explícita.",
        size=11, color=BODY)

# ============================================================ 4. AGENTES
heading("Catálogo de agentes especializados", "4")
body("Em vez de um único agente genérico, a empresa mantém agentes com escopos e "
     "permissões definidos. Cada um usa o harness loop e o roteamento multimodelo, "
     "e pode delegar a sub-agentes em paralelo.")

ag = doc.add_table(rows=8, cols=3)
ag.alignment = WD_TABLE_ALIGNMENT.CENTER
no_table_borders(ag)
arows = [
    ("Agente", "Função", "Conectores"),
    ("Assistente Executivo", "Triagem de e-mail, agenda, resumos diários, preparo de reuniões.", "Gmail, Calendar, Drive"),
    ("Atendimento / SDR", "Responde clientes no WhatsApp, qualifica leads, agenda demos.", "WhatsApp, sist. Comercial"),
    ("PMO / Projetos", "Atualiza status, gera relatórios e antecipa riscos nos projetos.", "sist. Projetos, Drive"),
    ("Contratos", "Gera, organiza e acompanha prazos de contratos; OCR de documentos.", "sist. Contratos, visão/OCR"),
    ("Suporte / IT Helpdesk", "Abre e resolve chamados, executa runbooks, opera máquinas.", "computer-use, ITSM"),
    ("DevOps / Engenharia", "Revisa PRs, corrige CI, evolui os sistemas internos.", "GitHub, computer-use"),
    ("Financeiro", "Concilia, organiza despesas e gera relatórios a partir de notas.", "sist. Contratos, ERP"),
]
widths = [Inches(1.9), Inches(3.3), Inches(1.4)]
for i, row in enumerate(arows):
    for j, txt in enumerate(row):
        c = ag.rows[i].cells[j]; c.width = widths[j]; set_cell_margins(c)
        if i == 0:
            set_cell_bg(c, "0E6E92")
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            add_run(p, txt, size=10, color=WHITE, bold=True)
        else:
            bg = "F1F4F6" if i % 2 else "FFFFFF"
            set_cell_bg(c, bg)
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            add_run(p, txt, size=9.5, color=(INK if j == 0 else BODY),
                    bold=(j == 0))

doc.add_page_break()

# ============================================================ 5. SISTEMAS INTERNOS
heading("Sistemas internos existentes — integrar e fortalecer", "5")
body("A setup.com.br já desenvolveu, de forma ágil (“vibe coding”), sistemas "
     "internos que sustentam a operação: gestão de projetos, gestão de contratos "
     "e a área comercial. A nova infraestrutura não os substitui — ela os "
     "encapsula como conectores e os coloca no alcance dos agentes.")
subheading("Integrar — esses sistemas viram conectores dos agentes")
bullet("o agente de PMO lê e atualiza status, monta relatórios e antecipa riscos.", "Gestão de Projetos — ")
bullet("o agente de Contratos gera minutas, acompanha prazos e aciona renovações.", "Gestão de Contratos — ")
bullet("o agente de Atendimento/SDR qualifica leads e move o funil no próprio sistema comercial.", "Comercial — ")
subheading("Fortalecer — a IA também cuida da saúde desses sistemas")
body("Sistemas nascidos em vibe coding entregam valor rápido, mas tendem a "
     "acumular dívida técnica. O agente de DevOps ajuda a estabilizá-los — "
     "adicionando testes, revisando alterações e documentando — reduzindo "
     "retrabalho e tornando-os mais confiáveis ao longo do tempo. O investimento "
     "feito é preservado e ampliado, não descartado.")

# ============================================================ 6. PESSOAS / CLIMA
heading("Pessoas no centro — clima, segurança psicológica e condições de trabalho", "6")
body("Tecnologia só compensa se melhora a vida de quem trabalha. O maior efeito "
     "desta infraestrutura não é a automação em si, mas o que ela faz pelas "
     "pessoas: tira de cima delas o trabalho repetitivo e penoso e devolve tempo, "
     "energia e tranquilidade.")
add_image(f"{A}/06_pessoas.png", 6.6, "Figura 4 — Como a infraestrutura melhora as condições de trabalho.")

subheading("Melhores condições de trabalho")
bullet("a IA assume triagem de e-mail, atualização de status e montagem de documentos, reduzindo sobrecarga e horas extras.", "Menos trabalho braçal — ")
bullet("o mesmo assistente no WhatsApp e no celular reduz a fricção e apoia o equilíbrio entre vida e trabalho.", "Flexibilidade — ")
bullet("um mentor sempre disponível nivela o jogo: profissionais novos ou juniores entregam com mais confiança.", "Apoio constante — ")

subheading("Mais segurança psicológica")
body("Segurança psicológica é o ambiente em que as pessoas se sentem seguras "
     "para propor ideias, pedir ajuda e errar sem medo de exposição. A "
     "infraestrutura reforça isso de forma concreta:")
bullet("dá para perguntar à IA quantas vezes for preciso, sem constrangimento de “perguntar besteira”.", "Tirar dúvidas sem julgamento — ")
bullet("como tudo passa por rascunho revisável e aprovação humana, o custo de errar cai — erra-se no rascunho, não na frente do cliente.", "O erro fica no rascunho — ")
bullet("a IA aumenta as pessoas; ela não as vigia nem mede desempenho individual. O controle das decisões é sempre humano.", "Aumenta, não vigia — ")

# caixa princípio — clima
t = doc.add_table(rows=1, cols=1); no_table_borders(t)
cell = t.rows[0].cells[0]; set_cell_bg(cell, "E1F0F4"); set_cell_margins(cell, 140, 140, 200, 200)
p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
add_run(p, "Efeito no clima:  ", size=11, color=INK, bold=True)
add_run(p, "equipes menos sobrecarregadas e mais confiantes colaboram melhor, "
        "retêm talentos e atendem clientes com mais qualidade. A IA passa a ser "
        "vista como aliada do profissional — não como ameaça ao seu trabalho.",
        size=11, color=BODY)

doc.add_page_break()

# ============================================================ 7. GOVERNANÇA
heading("Segurança e governança — não-negociável", "7")
body("A produtividade só é sustentável com confiança. A infraestrutura aplica "
     "controle por padrão:")
bullet("login via Google Workspace corporativo; cada agente age como o usuário, herdando apenas as permissões que ele já tem.", "Identidade única (SSO) — ")
bullet("cada conector recebe somente os escopos necessários (ex.: Drive apenas nas pastas autorizadas).", "Permissões mínimas — ")
bullet("envios a clientes, exclusões e alterações irreversíveis passam por confirmação humana.", "Aprovação de ações externas — ")
bullet("chaves e tokens em cofre (Vault / Secret Manager), nunca no código.", "Segredos protegidos — ")
bullet("cada passo do loop (ferramenta, parâmetros, resultado, aprovação) é registrado para conformidade.", "Auditoria completa — ")
bullet("dados da empresa não treinam modelos públicos; ambientes com garantia de não retenção.", "Isolamento de dados — ")

# ============================================================ 8. BASE EXISTENTE
heading("O que já existe e o que falta construir", "8")
body("Boa parte da fundação já está no ambiente atual da setup.com.br, o que "
     "reduz o tempo e o risco de implantação:")
subheading("Já disponível")
bullet("conector de WhatsApp (número pessoal via automação de navegador e API oficial para produção).", "WhatsApp — ")
bullet("Gmail, Calendar e Drive já integrados como conectores MCP.", "Google Workspace — ")
bullet("gestão de projetos, contratos e comercial, prontos para serem expostos como conectores.", "Sistemas internos — ")
bullet("capacidade de operar o computador e integração com GitHub para as equipes técnicas.", "computer-use e GitHub — ")
subheading("A construir")
bullet("roteador multimodelo, catálogo de agentes, memória de longo prazo (base de conhecimento) e a camada de governança/SSO corporativa.", "")

# ============================================================ 9. ROADMAP
heading("Roadmap de implantação", "9")
add_image(f"{A}/04_roadmap.png", 6.6, "Figura 5 — Implantação faseada, do alicerce à otimização contínua.")

doc.add_page_break()

# ============================================================ 10. INDICADORES
heading("Indicadores de acompanhamento", "10")
body("O sucesso é medido por produtividade, qualidade e — em pé de igualdade — "
     "pelo bem-estar das equipes. Em vez de prometer números, definimos o que "
     "acompanhar: a linha de base é medida na própria operação e as metas saem "
     "dela. Assim evitamos projeções artificiais e ganhamos comparações honestas "
     "(antes e depois) sobre dados reais da setup.com.br.")
add_image(f"{A}/05_indicadores.png", 6.4, "Figura 6 — Domínios de indicadores; valores definidos a partir da medição inicial.")

# fechamento
sp = doc.add_paragraph(); sp.paragraph_format.space_before = Pt(10)
t = doc.add_table(rows=1, cols=1); no_table_borders(t)
cell = t.rows[0].cells[0]; set_cell_bg(cell, "003C54"); set_cell_margins(cell, 160, 160, 220, 220)
p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(0)
add_run(p, "Em uma frase:  ", size=12, color=TEAL, bold=True)
add_run(p, "um cérebro de IA, muitos corpos — a mesma inteligência em cada tela, "
        "operando com segurança para que cada profissional da setup.com.br "
        "entregue mais e melhor.", size=12, color=WHITE, bold=True)

# ---- rodapé com logo + paginação ----
def add_footer():
    for section in doc.sections:
        footer = section.footer
        # logo da marca (alinhado à esquerda, parágrafo próprio)
        lp = footer.paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        lp.paragraph_format.space_after = Pt(0)
        lp.add_run().add_picture(f"{A}/setup_logo@hi.png", height=Inches(0.16))
        # linha de texto + número de página
        p = footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(p, "setup.com.br · Infraestrutura de TI com IA · Documento Executivo · ",
                size=8, color=SLATE)
        # campo de número de página
        fldSimple = OxmlElement("w:fldSimple")
        fldSimple.set(qn("w:instr"), "PAGE")
        r = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "16"); rPr.append(sz)
        r.append(rPr)
        t_ = OxmlElement("w:t"); t_.text = "1"; r.append(t_)
        fldSimple.append(r)
        p._p.append(fldSimple)

add_footer()

doc.save(OUT)
print("Documento salvo em", OUT)
